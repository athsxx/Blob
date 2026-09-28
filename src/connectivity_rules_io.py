"""
Normalize, validate, merge, diff, and save connectivity_rules.json.

The DALIA inspection workbook uses Insert / Input View / Open To Check /
Output View. A hyphenated face such as C-F is one hole visible on both
cameras. Both the type-in editor and the spreadsheet path write the same
JSON contract that LogicEngine.load_rules already understands.
"""
from __future__ import annotations

import csv
import json
import os
import re
import shutil
import tempfile
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Sequence, Set, Tuple

from logic_engine import parse_face_to_faces, rule_has_unavailable_output

VALID_FACES = frozenset("ABCDEF")
DEFAULT_LOGIC = "AND"
DEFAULT_TIMING = {"max_delay_ms": 250, "min_stable_frames": 3}
OUTPUT_FACE_CHOICES = [
    "A", "B", "C", "D", "E", "F",
    "C_F", "A_F", "A_E", "A___E", "B_C", "B_OR_D",
]

# After spaces/hyphens → underscore. Keys are uppercased alias keys.
_HOLE_ALIASES = {
    "SEC_P_P": "SEC_P_P",
    "P_P": "SEC_P_P",
    "SEC_M_M": "SEC_M_M",
    "M_M": "SEC_M_M",
    "SEC_N_N": "SEC_N_N",
    "N_N": "SEC_N_N",
    "SEC_Z_Z": "SEC_Z_Z",
    "Z_Z": "SEC_Z_Z",
    "SEC_T_T": "SEC_T_T",
    "T_T": "SEC_T_T",
    "SEC_JK_JK": "SEC_JK_JK",
    "JK_JK": "SEC_JK_JK",
    "CENTRE_HOLE": "CENTRE_HOLE",
    "CENTER_HOLE": "CENTRE_HOLE",
}

_HEADER_ALIASES = {
    "input_face": {
        "input face", "from face", "laser face", "laser in face",
        "input_face", "face", "a-f",
        "input view (flashlight)", "input view",
    },
    "input_hole": {
        "input hole", "from hole", "laser hole", "laser in hole",
        "input_hole", "hole", "section", "hole id", "input hole id",
        "insert rod / passlight", "insert rod", "passlight",
    },
    "output_face": {
        "output face", "to face", "exit face", "light face",
        "output_face", "light should appear face",
        "output view (lightbulb)", "output view",
    },
    "output_hole": {
        "output hole", "to hole", "exit hole", "light hole",
        "output_hole", "light should appear hole", "output hole id",
        "open to check (for inspection)", "open to check",
    },
    "logic": {"logic", "and/or", "and or"},
    "mandatory": {"mandatory", "required"},
}

_OUTPUT_N_FACE = re.compile(r"^output\s*([0-9]+)\s*face$")
_OUTPUT_N_HOLE = re.compile(r"^output\s*([0-9]+)\s*hole$")


@dataclass
class ConnectionRow:
    input_face: str
    input_hole: str
    output_face: str
    output_hole: str
    mandatory: bool = True
    logic: str = DEFAULT_LOGIC
    source_row: Optional[int] = None


@dataclass
class SpreadsheetParse:
    headers: List[str] = field(default_factory=list)
    mapped: Dict[str, int] = field(default_factory=dict)
    extra_outputs: List[Tuple[int, int, int]] = field(default_factory=list)  # (n, face_col, hole_col)
    rows: List[ConnectionRow] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    skipped_title_rows: int = 0
    source_row_count: int = 0


@dataclass
class RulesDiff:
    added: List[str]
    removed: List[str]
    changed: List[str]
    unchanged: int
    merged_from_rows: int
    duplicate_ids_in_old: List[str]
    unknown_holes: List[Tuple[str, str]]
    guided_before: int
    guided_after: int
    old_count: int
    new_count: int

    def is_noop(self) -> bool:
        return not self.added and not self.removed and not self.changed


def _nfc(text: str) -> str:
    return unicodedata.normalize("NFC", str(text or ""))


def normalize_hole_id(raw: Any) -> str:
    text = _nfc(raw).strip()
    if not text:
        return ""
    text = text.replace("ø", "Ø").replace("Ø".lower(), "Ø")
    text = text.replace("phi", "Ø").replace("PHI", "Ø")
    alias_key = re.sub(r"[^A-Z0-9Ø]+", "_", text.upper())
    alias_key = re.sub(r"_+", "_", alias_key).strip("_")
    if alias_key in _HOLE_ALIASES:
        return _HOLE_ALIASES[alias_key]
    return alias_key


def is_unnamed_counter(raw: Any) -> bool:
    """T-T and Z-Z counters are two holes each and are not named yet."""
    key = normalize_hole_id(raw)
    return "COUNTER" in key and ("T_T" in key or "Z_Z" in key)


def normalize_input_face(raw: Any) -> str:
    """A-F / C-F is one hole visible on both cameras, same form as an output face."""
    return normalize_output_face(raw)


def normalize_output_face(raw: Any) -> str:
    text = _nfc(raw).strip().upper()
    if not text:
        return ""
    compact = re.sub(r"\s+", " ", text)
    if compact in OUTPUT_FACE_CHOICES or compact in VALID_FACES:
        return compact
    if " OR " in f" {compact} " or "/OR/" in compact.replace(" ", ""):
        letters = [ch for ch in compact if ch in VALID_FACES]
        if len(letters) >= 2:
            return f"{letters[0]}_OR_{letters[1]}"
    sep = compact.replace(" ", "").replace("/", "-").replace("_OR_", "-")
    if "-" in sep:
        letters = [p for p in sep.split("-") if p in VALID_FACES]
        if len(letters) == 1:
            return letters[0]
        if len(letters) >= 2:
            known = f"{letters[0]}_{letters[1]}"
            or_form = f"{letters[0]}_OR_{letters[1]}"
            if known in OUTPUT_FACE_CHOICES:
                return known
            if or_form in OUTPUT_FACE_CHOICES:
                return or_form
            return known
    letters = [ch for ch in compact if ch in VALID_FACES]
    if len(letters) == 1:
        return letters[0]
    parsed = parse_face_to_faces(compact.replace("-", "_").replace("/", "_"))
    if parsed:
        if compact.replace(" ", "") in OUTPUT_FACE_CHOICES:
            return compact.replace(" ", "")
        if len(parsed) == 1:
            return parsed[0]
        return "_".join(parsed)
    return compact.replace("-", "_").replace("/", "_").replace(" ", "_")


def make_rule_id(face: str, hole_id: str) -> str:
    return f"FACE_{face}_{hole_id}"


def split_cell_list(raw: Any) -> List[str]:
    """Split a cell on commas that are not inside parentheses."""
    text = _nfc(raw).strip()
    if not text:
        return []
    parts: List[str] = []
    buf: List[str] = []
    depth = 0
    for ch in text:
        if ch == "(":
            depth += 1
            buf.append(ch)
        elif ch == ")":
            depth = max(0, depth - 1)
            buf.append(ch)
        elif ch == "," and depth == 0:
            part = "".join(buf).strip()
            if part:
                parts.append(part)
            buf = []
        else:
            buf.append(ch)
    part = "".join(buf).strip()
    if part:
        parts.append(part)
    return parts


def _join_faces(faces: Sequence[str]) -> str:
    letters: List[str] = []
    for face in faces:
        for letter in parse_face_to_faces(normalize_output_face(face)) or []:
            if letter not in letters:
                letters.append(letter)
    if not letters:
        return ""
    if len(letters) == 1:
        return letters[0]
    return "-".join(letters)


def pair_exit_holes_and_faces(
    holes: Sequence[str], faces: Sequence[str]
) -> List[Tuple[str, str]]:
    """Pair exits from the left.

    Extra holes stay on the last face. Extra faces mean the last hole is
    visible on each of those faces.
    """
    hole_list = [h for h in holes if str(h).strip()]
    face_list = [f for f in faces if str(f).strip()]
    if not hole_list or not face_list:
        return []
    if len(hole_list) == len(face_list):
        return list(zip(hole_list, face_list))
    if len(hole_list) > len(face_list):
        head = len(face_list) - 1
        paired = list(zip(hole_list[:head], face_list[:head]))
        last_face = face_list[-1]
        paired.extend((hole, last_face) for hole in hole_list[head:])
        return paired
    head = len(hole_list) - 1
    paired = list(zip(hole_list[:head], face_list[:head]))
    paired.append((hole_list[-1], _join_faces(face_list[head:])))
    return paired


def _truthy(raw: Any) -> bool:
    if raw is None or raw is True:
        return True
    if raw is False:
        return False
    text = str(raw).strip().lower()
    if text in {"", "1", "true", "yes", "y", "mandatory", "required"}:
        return True
    if text in {"0", "false", "no", "n", "optional"}:
        return False
    return True


def _logic(raw: Any) -> str:
    text = str(raw or DEFAULT_LOGIC).strip().upper()
    return "OR" if text == "OR" else DEFAULT_LOGIC


def canonical_rule(rule: Dict[str, Any]) -> Dict[str, Any]:
    inp = rule.get("input") or {}
    face = normalize_input_face(inp.get("face", ""))
    hole = normalize_hole_id(inp.get("hole_id", ""))
    outputs: List[Dict[str, Any]] = []
    seen = set()
    for out in rule.get("expected_outputs") or []:
        oface = normalize_output_face(out.get("face", ""))
        ohole = normalize_hole_id(out.get("hole_id", ""))
        key = (oface, ohole)
        if not oface or not ohole or key in seen:
            continue
        seen.add(key)
        outputs.append({
            "face": oface,
            "hole_id": ohole,
            "mandatory": bool(out.get("mandatory", True)),
        })
    timing = rule.get("timing") or {}
    return {
        "rule_id": str(rule.get("rule_id") or make_rule_id(face, hole)),
        "input": {"face": face, "hole_id": hole},
        "expected_outputs": outputs,
        "logic": _logic(rule.get("logic")),
        "timing": {
            "max_delay_ms": int(timing.get("max_delay_ms") or DEFAULT_TIMING["max_delay_ms"]),
            "min_stable_frames": int(
                timing.get("min_stable_frames") or DEFAULT_TIMING["min_stable_frames"]
            ),
        },
    }


def merge_connection_rows(rows: Sequence[ConnectionRow]) -> List[Dict[str, Any]]:
    grouped: Dict[Tuple[str, str], Dict[str, Any]] = {}
    order: List[Tuple[str, str]] = []
    for row in rows:
        face = normalize_input_face(row.input_face)
        hole = normalize_hole_id(row.input_hole)
        oface = normalize_output_face(row.output_face)
        ohole = normalize_hole_id(row.output_hole)
        if not face or not hole or not oface or not ohole:
            continue
        key = (face, hole)
        if key not in grouped:
            grouped[key] = {
                "rule_id": make_rule_id(face, hole),
                "input": {"face": face, "hole_id": hole},
                "expected_outputs": [],
                "logic": _logic(row.logic),
                "timing": dict(DEFAULT_TIMING),
            }
            order.append(key)
        outs = grouped[key]["expected_outputs"]
        seen = {(o["face"], o["hole_id"]) for o in outs}
        if (oface, ohole) not in seen:
            outs.append({
                "face": oface,
                "hole_id": ohole,
                "mandatory": bool(row.mandatory),
            })
        if row.logic:
            grouped[key]["logic"] = _logic(row.logic)
    rules = [grouped[k] for k in order]
    return uniquify_rule_ids(rules)


def rules_to_connection_rows(rules: Sequence[Dict[str, Any]]) -> List[ConnectionRow]:
    rows: List[ConnectionRow] = []
    for rule in rules or []:
        canon = canonical_rule(rule)
        inp = canon["input"]
        logic = canon["logic"]
        outputs = canon["expected_outputs"] or [
            {"face": "", "hole_id": "", "mandatory": True}
        ]
        for out in outputs:
            rows.append(ConnectionRow(
                input_face=inp["face"],
                input_hole=inp["hole_id"],
                output_face=out.get("face", ""),
                output_hole=out.get("hole_id", ""),
                mandatory=bool(out.get("mandatory", True)),
                logic=logic,
            ))
    return rows


def uniquify_rule_ids(rules: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """One rule_id per rule. Same input hole already merged; suffix only on true collisions."""
    used: Dict[str, int] = {}
    out: List[Dict[str, Any]] = []
    for rule in rules:
        canon = canonical_rule(rule)
        rid = canon["rule_id"] or make_rule_id(
            canon["input"]["face"], canon["input"]["hole_id"]
        )
        n = used.get(rid, 0)
        if n:
            canon["rule_id"] = f"{rid}_{n + 1}"
        used[rid] = n + 1
        out.append(canon)
    return out


def duplicate_rule_ids(rules: Sequence[Dict[str, Any]]) -> List[str]:
    counts: Dict[str, int] = {}
    for rule in rules or []:
        rid = str(rule.get("rule_id") or "")
        if rid:
            counts[rid] = counts.get(rid, 0) + 1
    return [rid for rid, n in counts.items() if n > 1]


def validate_rules(rules: Sequence[Dict[str, Any]]) -> List[str]:
    errors: List[str] = []
    seen_ids: Set[str] = set()
    for i, raw in enumerate(rules or [], start=1):
        rule = canonical_rule(raw)
        rid = rule["rule_id"]
        face = rule["input"]["face"]
        hole = rule["input"]["hole_id"]
        if not parse_face_to_faces(face):
            errors.append(f"Rule {i} ({rid}): laser face must be A–F, got {face!r}")
        if not hole:
            errors.append(f"Rule {i} ({rid}): laser hole is empty")
        if rid in seen_ids:
            errors.append(f"Rule {i}: duplicate id {rid}")
        seen_ids.add(rid)
        if not rule["expected_outputs"]:
            errors.append(f"Rule {i} ({rid}): no light-out holes")
        for out in rule["expected_outputs"]:
            parsed = parse_face_to_faces(out["face"])
            if not parsed:
                errors.append(
                    f"Rule {i} ({rid}): cannot read light face {out['face']!r}"
                )
            if not out["hole_id"]:
                errors.append(f"Rule {i} ({rid}): empty light hole")
        if rule["logic"] not in {"AND", "OR"}:
            errors.append(f"Rule {i} ({rid}): logic must be AND or OR")
    return errors


def hole_ids_for_face(rules: Sequence[Dict[str, Any]], face: str) -> List[str]:
    """Unique hole_ids that appear as input or expected output on this face."""
    want = str(face).strip().upper()
    names: List[str] = []
    seen: Set[str] = set()
    for rule in rules or []:
        inp = rule.get("input") or {}
        if want in parse_face_to_faces(str(inp.get("face", ""))):
            hid = str(inp.get("hole_id") or "").strip()
            if hid and hid not in seen:
                seen.add(hid)
                names.append(hid)
        for out in rule.get("expected_outputs") or []:
            faces = parse_face_to_faces(str(out.get("face", "")))
            if want not in faces:
                continue
            hid = str(out.get("hole_id") or "").strip()
            if hid and hid not in seen:
                seen.add(hid)
                names.append(hid)
    return names


def roi_names_by_face(config_dir: str, folder: str) -> Dict[str, Set[str]]:
    from config_loader import face_roi_basename

    names: Dict[str, Set[str]] = {f: set() for f in "ABCDEF"}
    for face in "ABCDEF":
        path = os.path.join(config_dir, folder, face_roi_basename(face))
        if not os.path.isfile(path):
            continue
        try:
            with open(path, "r", encoding="utf-8") as f:
                data = json.load(f)
        except (OSError, json.JSONDecodeError):
            continue
        for circle in data.get("circles") or []:
            hid = str(circle.get("name") or circle.get("hole_id") or "").strip()
            if hid:
                names[face].add(hid)
    return names


def unknown_holes_vs_rois(
    rules: Sequence[Dict[str, Any]],
    roi_names: Dict[str, Set[str]],
) -> List[Tuple[str, str]]:
    missing: List[Tuple[str, str]] = []
    seen: Set[Tuple[str, str]] = set()
    known_rules = {f: set(hole_ids_for_face(rules, f)) for f in "ABCDEF"}
    for face in "ABCDEF":
        for hid in sorted(known_rules[face]):
            rois = roi_names.get(face) or set()
            if hid in rois:
                continue
            # Excel-normalized name vs a spaced ROI label
            aliases = {hid, hid.replace("_", " "), hid.replace("_", "-")}
            if any(a in rois for a in aliases):
                continue
            key = (face, hid)
            if key not in seen:
                seen.add(key)
                missing.append(key)
    return missing


def guided_step_count(
    rules: Sequence[Dict[str, Any]],
    available_faces: Optional[Set[str]] = None,
) -> int:
    faces = available_faces if available_faces is not None else set("ABCDEF")
    seen: Set[str] = set()
    n = 0
    for rule in rules or []:
        rid = str(rule.get("rule_id") or "")
        if not rid or rid in seen:
            continue
        if rule_has_unavailable_output(rule, faces):
            continue
        input_faces = parse_face_to_faces(str((rule.get("input") or {}).get("face", "")))
        if input_faces and all(f not in faces for f in input_faces):
            continue
        seen.add(rid)
        n += 1
    return n


def _input_key(rule: Dict[str, Any]) -> Tuple[str, str]:
    inp = (rule.get("input") or {})
    return (str(inp.get("face") or ""), str(inp.get("hole_id") or ""))


def _index_by_input(rules: Sequence[Dict[str, Any]]) -> Dict[Tuple[str, str], Dict[str, Any]]:
    """One entry per laser-in hole; later duplicate rule_ids merge their outputs."""
    indexed: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for rule in rules or []:
        canon = canonical_rule(rule)
        key = _input_key(canon)
        if not key[0] or not key[1]:
            continue
        if key not in indexed:
            indexed[key] = canon
            continue
        existing = indexed[key]
        seen = {(o["face"], o["hole_id"]) for o in existing["expected_outputs"]}
        for out in canon["expected_outputs"]:
            t = (out["face"], out["hole_id"])
            if t not in seen:
                existing["expected_outputs"].append(out)
                seen.add(t)
    return indexed


def _outputs_key(rule: Dict[str, Any]) -> Tuple:
    outs = tuple(sorted(
        (o["face"], o["hole_id"], bool(o.get("mandatory", True)))
        for o in rule.get("expected_outputs") or []
    ))
    return (rule.get("logic"), outs)


def diff_rules(
    old_rules: Sequence[Dict[str, Any]],
    new_rules: Sequence[Dict[str, Any]],
    merged_from_rows: int = 0,
    roi_names: Optional[Dict[str, Set[str]]] = None,
    available_faces: Optional[Set[str]] = None,
) -> RulesDiff:
    old_idx = _index_by_input(old_rules)
    new_idx = _index_by_input(new_rules)
    added = sorted(
        new_idx[k]["rule_id"] for k in new_idx if k not in old_idx
    )
    removed = sorted(
        old_idx[k]["rule_id"] for k in old_idx if k not in new_idx
    )
    changed = sorted(
        new_idx[k]["rule_id"]
        for k in new_idx
        if k in old_idx and _outputs_key(new_idx[k]) != _outputs_key(old_idx[k])
    )
    unknown = unknown_holes_vs_rois(new_rules, roi_names or {})
    return RulesDiff(
        added=added,
        removed=removed,
        changed=changed,
        unchanged=len(new_idx) - len(added) - len(changed),
        merged_from_rows=merged_from_rows,
        duplicate_ids_in_old=duplicate_rule_ids(old_rules),
        unknown_holes=unknown,
        guided_before=guided_step_count(old_rules, available_faces),
        guided_after=guided_step_count(new_rules, available_faces),
        old_count=len(list(old_rules or [])),
        new_count=len(new_idx),
    )


def format_diff(diff: RulesDiff) -> str:
    lines = [
        f"This replaces {diff.old_count} working rules with {diff.new_count} unique connections.",
        f"Guided steps now: {diff.guided_after} (was {diff.guided_before}).",
    ]
    if diff.merged_from_rows:
        lines.append(
            f"Spreadsheet rows merged: {diff.merged_from_rows} → {diff.new_count} rules."
        )
    if diff.duplicate_ids_in_old:
        lines.append(
            "Duplicate ids in the current file that merge on save: "
            + ", ".join(diff.duplicate_ids_in_old)
        )
    lines.append(f"Added ({len(diff.added)}): " + (", ".join(diff.added) or "none"))
    lines.append(f"Removed ({len(diff.removed)}): " + (", ".join(diff.removed) or "none"))
    lines.append(f"Changed ({len(diff.changed)}): " + (", ".join(diff.changed) or "none"))
    if diff.unknown_holes:
        lines.append("")
        lines.append("Hole names with no matching ellipse on the face picture:")
        for face, hid in diff.unknown_holes[:40]:
            lines.append(f"  Face {face}: {hid}")
        if len(diff.unknown_holes) > 40:
            lines.append(f"  … {len(diff.unknown_holes) - 40} more")
        lines.append(
            "Those names must match the labels on the face pictures. "
            "If they do not match, every check will fail. "
            "Place missing ellipses with Place hole ROIs on the camera setup page."
        )
    lines.append("")
    lines.append("Timing in the file is stored only. PASS/FAIL still uses the 2 second window.")
    return "\n".join(lines)


def save_rules_atomic(path: str, rules: Sequence[Dict[str, Any]], backup: bool = True) -> str:
    errors = validate_rules(rules)
    if errors:
        raise ValueError("\n".join(errors))
    payload = {"rules": uniquify_rule_ids(rules)}
    path = os.path.abspath(path)
    directory = os.path.dirname(path)
    os.makedirs(directory, exist_ok=True)
    if backup and os.path.isfile(path):
        shutil.copy2(path, path + ".bak")
    fd, tmp = tempfile.mkstemp(prefix="rules_", suffix=".json", dir=directory)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2, ensure_ascii=False)
            f.write("\n")
        os.replace(tmp, path)
    except Exception:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise
    return path


def load_rules_file(path: str) -> List[Dict[str, Any]]:
    if not os.path.isfile(path):
        return []
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    rules = data.get("rules") or []
    return list(rules) if isinstance(rules, list) else []


def _norm_header(text: Any) -> str:
    return re.sub(r"\s+", " ", _nfc(text).strip().lower())


def _map_headers(headers: Sequence[str]) -> Tuple[Dict[str, int], List[Tuple[int, int, int]]]:
    mapped: Dict[str, int] = {}
    extras: Dict[int, Dict[str, int]] = {}
    for i, raw in enumerate(headers):
        h = _norm_header(raw)
        if not h:
            continue
        for field_name, aliases in _HEADER_ALIASES.items():
            if h in aliases and field_name not in mapped:
                mapped[field_name] = i
                break
        mface = _OUTPUT_N_FACE.match(h)
        mhole = _OUTPUT_N_HOLE.match(h)
        if mface:
            extras.setdefault(int(mface.group(1)), {})["face"] = i
        if mhole:
            extras.setdefault(int(mhole.group(1)), {})["hole"] = i
    extra_pairs = [
        (n, cols["face"], cols["hole"])
        for n, cols in sorted(extras.items())
        if "face" in cols and "hole" in cols
    ]
    return mapped, extra_pairs


def _cell(row: Sequence[Any], index: Optional[int]) -> str:
    if index is None or index < 0 or index >= len(row):
        return ""
    val = row[index]
    if val is None:
        return ""
    return str(val).strip()


def _rows_from_table(headers: Sequence[str], data_rows: Sequence[Sequence[Any]]) -> SpreadsheetParse:
    parsed = SpreadsheetParse(headers=[str(h) for h in headers])
    mapped, extras = _map_headers(headers)
    parsed.mapped = mapped
    parsed.extra_outputs = extras
    needed = {"input_face", "input_hole", "output_face", "output_hole"}
    if not needed.issubset(mapped):
        missing = sorted(needed - set(mapped))
        parsed.errors.append(
            "Could not map spreadsheet columns ("
            + ", ".join(missing)
            + "). Headers found: "
            + ", ".join(h for h in parsed.headers if str(h).strip())
            + ". Share the workbook so these headers can be bound."
        )
        return parsed
    for n, row in enumerate(data_rows, start=1):
        if not any(_cell(row, i) for i in range(len(row))):
            continue
        parsed.source_row_count += 1
        base = ConnectionRow(
            input_face=_cell(row, mapped["input_face"]),
            input_hole=_cell(row, mapped["input_hole"]),
            output_face=_cell(row, mapped["output_face"]),
            output_hole=_cell(row, mapped["output_hole"]),
            mandatory=_truthy(_cell(row, mapped.get("mandatory"))) if "mandatory" in mapped else True,
            logic=_logic(_cell(row, mapped.get("logic"))) if "logic" in mapped else DEFAULT_LOGIC,
            source_row=n,
        )
        if not base.input_hole and not base.output_hole:
            continue
        inserts = split_cell_list(base.input_hole)
        exits = pair_exit_holes_and_faces(
            split_cell_list(base.output_hole),
            split_cell_list(base.output_face),
        )
        if len(inserts) != 1 or not exits:
            parsed.errors.append(
                f"Row {n}: held ({base.input_hole or 'blank insert'} → {base.output_hole})"
            )
            continue
        if is_unnamed_counter(inserts[0]) or any(is_unnamed_counter(hole) for hole, _face in exits):
            parsed.errors.append(
                f"Row {n}: held, unnamed counter ({inserts[0]})"
            )
            continue
        base.input_hole = inserts[0]
        parsed.rows.append(ConnectionRow(
            input_face=base.input_face,
            input_hole=base.input_hole,
            output_face=exits[0][1],
            output_hole=exits[0][0],
            mandatory=base.mandatory,
            logic=base.logic,
            source_row=n,
        ))
        for hole, face in exits[1:]:
            parsed.rows.append(ConnectionRow(
                input_face=base.input_face,
                input_hole=base.input_hole,
                output_face=face,
                output_hole=hole,
                mandatory=base.mandatory,
                logic=base.logic,
                source_row=n,
            ))
        for _n, fcol, hcol in extras:
            oface = _cell(row, fcol)
            ohole = _cell(row, hcol)
            if oface and ohole:
                parsed.rows.append(ConnectionRow(
                    input_face=base.input_face,
                    input_hole=base.input_hole,
                    output_face=oface,
                    output_hole=ohole,
                    mandatory=base.mandatory,
                    logic=base.logic,
                    source_row=n,
                ))
    return parsed


def _find_header_row(rows: Sequence[Sequence[Any]]) -> Tuple[int, List[str]]:
    best_i = 0
    best_score = -1
    best_headers: List[str] = []
    for i, row in enumerate(rows[:15]):
        headers = [str(c) if c is not None else "" for c in row]
        mapped, extras = _map_headers(headers)
        score = len(mapped) + len(extras)
        if score > best_score:
            best_score = score
            best_i = i
            best_headers = headers
    return best_i, best_headers


def parse_csv_file(path: str) -> SpreadsheetParse:
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = [list(r) for r in reader]
    if not rows:
        parsed = SpreadsheetParse()
        parsed.errors.append("Spreadsheet is empty.")
        return parsed
    header_i, headers = _find_header_row(rows)
    parsed = _rows_from_table(headers, rows[header_i + 1 :])
    parsed.skipped_title_rows = header_i
    return parsed


def parse_xlsx_file(path: str) -> SpreadsheetParse:
    try:
        import openpyxl
    except ImportError as exc:
        parsed = SpreadsheetParse()
        parsed.errors.append(
            "openpyxl is not installed. pip install openpyxl  — or Save As CSV and load that."
        )
        parsed.errors.append(str(exc))
        return parsed
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() in {"rules", "connectivity", "hole connections", "connections"}:
            sheet = wb[name]
            break
    if sheet is None:
        sheet = wb.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([c if c is not None else "" for c in row])
    wb.close()
    if not rows:
        parsed = SpreadsheetParse()
        parsed.errors.append("Spreadsheet is empty.")
        return parsed
    header_i, headers = _find_header_row(rows)
    parsed = _rows_from_table(headers, rows[header_i + 1 :])
    parsed.skipped_title_rows = header_i
    return parsed


def parse_spreadsheet(path: str) -> SpreadsheetParse:
    ext = os.path.splitext(path)[1].lower()
    if ext in {".xlsx", ".xlsm"}:
        return parse_xlsx_file(path)
    if ext in {".csv", ".txt"}:
        return parse_csv_file(path)
    parsed = SpreadsheetParse()
    parsed.errors.append(f"Unsupported file type {ext}. Use .xlsx or .csv.")
    return parsed


def without_conflicting_inputs(
    rows: Sequence[ConnectionRow],
) -> Tuple[List[ConnectionRow], List[str]]:
    """Drop an insert that the sheet describes twice with different exits."""
    by_key: Dict[Tuple[str, str], Dict[int, Set[Tuple[str, str]]]] = {}
    for row in rows:
        key = (normalize_input_face(row.input_face), normalize_hole_id(row.input_hole))
        src = row.source_row if row.source_row is not None else id(row)
        by_key.setdefault(key, {}).setdefault(src, set()).add(
            (normalize_output_face(row.output_face), normalize_hole_id(row.output_hole))
        )
    blocked: Set[Tuple[str, str]] = set()
    notes: List[str] = []
    for key, sources in by_key.items():
        signatures = {frozenset(exits) for exits in sources.values()}
        if len(signatures) > 1:
            blocked.add(key)
            notes.append(
                f"Held overlapping insert {key[0]} {key[1]} "
                f"({len(signatures)} different exit lists)"
            )
    kept = [
        row for row in rows
        if (normalize_input_face(row.input_face), normalize_hole_id(row.input_hole)) not in blocked
    ]
    return kept, notes


def rules_from_spreadsheet(path: str) -> Tuple[List[Dict[str, Any]], SpreadsheetParse]:
    parsed = parse_spreadsheet(path)
    if parsed.errors and not parsed.rows:
        return [], parsed
    kept, notes = without_conflicting_inputs(parsed.rows)
    parsed.rows = kept
    parsed.errors.extend(notes)
    rules = merge_connection_rows(kept)
    return rules, parsed
